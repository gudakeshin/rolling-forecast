"""Board pack export — PowerPoint and PDF from published forecast versions."""

from __future__ import annotations

import io
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from app.models.forecast import ForecastVersion, ForecastLineResult, ModelMetadata
from app.models.line_item import LineItem

logger = logging.getLogger(__name__)


def _aggregate_by_category(db: Session, version_id: str) -> list[dict[str, Any]]:
    results = (
        db.query(ForecastLineResult, LineItem)
        .join(LineItem, LineItem.id == ForecastLineResult.line_item_id)
        .filter(ForecastLineResult.version_id == version_id)
        .all()
    )
    buckets: dict[str, float] = {}
    for r, li in results:
        val = r.override_value if r.is_overridden else r.p50
        buckets[li.category] = buckets.get(li.category, 0.0) + float(val or 0)
    return [{"category": k, "total": v} for k, v in sorted(buckets.items())]


def _exog_footnotes(db: Session, version_id: str) -> list[str]:
    """Lines whose published point rests on admitted exogenous drivers."""
    rows = (
        db.query(ForecastLineResult, LineItem, ModelMetadata)
        .join(LineItem, LineItem.id == ForecastLineResult.line_item_id)
        .outerjoin(ModelMetadata, ModelMetadata.line_result_id == ForecastLineResult.id)
        .filter(ForecastLineResult.version_id == version_id)
        .all()
    )
    seen: set[int] = set()
    footnotes: list[str] = []
    for flr, li, meta in rows:
        if li.id in seen:
            continue
        params = (meta.parameters if meta else None) or {}
        spec = params.get("exog_spec") or {}
        if not spec.get("exog_admitted"):
            continue
        seen.add(li.id)
        drivers = spec.get("drivers") or []
        labels = []
        for d in drivers:
            key = d.get("driver_key") or d.get("driver_id")
            lag = d.get("lag", 0)
            labels.append(f"{key}@lag{lag}")
        tag = ", ".join(labels) if labels else "exog"
        inflated = (
            " (intervals inflated for driver uncertainty)"
            if spec.get("exog_variance_inflated") or params.get("exog_variance_inflated")
            else ""
        )
        footnotes.append(f"{li.name} [{li.account_code}]: {tag}{inflated}")
    return footnotes


def build_board_pack_pptx(db: Session, version: ForecastVersion) -> bytes:
    """Generate a simple board-ready PowerPoint deck."""
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = version.name
    slide.placeholders[1].text = (
        f"Status: {version.status}\n"
        f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}\n"
        f"Horizon: {version.horizon_months} months"
    )

    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Forecast Summary by Category"
    body = slide2.placeholders[1].text_frame
    body.clear()
    for row in _aggregate_by_category(db, version.id):
        p = body.add_paragraph()
        p.text = f"{row['category']}: ${row['total']:,.0f}"
        p.level = 0

    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    slide3.shapes.title.text = "Confidence Distribution"
    body3 = slide3.placeholders[1].text_frame
    body3.clear()
    for label, count in [
        ("High", version.high_confidence_count),
        ("Medium", version.medium_confidence_count),
        ("Low", version.low_confidence_count),
        ("Overrides", version.override_count),
    ]:
        p = body3.add_paragraph()
        p.text = f"{label}: {count}"
        p.level = 0

    footnotes = _exog_footnotes(db, version.id)
    if footnotes:
        slide4 = prs.slides.add_slide(prs.slide_layouts[1])
        slide4.shapes.title.text = "Driver-dependent figures"
        body4 = slide4.placeholders[1].text_frame
        body4.clear()
        note = body4.paragraphs[0]
        note.text = (
            "These lines use exogenous drivers; disabling driver forecasting "
            "does not un-publish them."
        )
        for line in footnotes[:40]:
            p = body4.add_paragraph()
            p.text = line
            p.level = 0

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_board_pack_pdf(db: Session, version: ForecastVersion) -> bytes:
    """Generate a simple board-ready PDF summary."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=letter)
        width, height = letter
        y = height - 72
        c.setFont("Helvetica-Bold", 18)
        c.drawString(72, y, f"Board Pack: {version.name}")
        y -= 28
        c.setFont("Helvetica", 11)
        c.drawString(72, y, f"Status: {version.status}  |  Horizon: {version.horizon_months} months")
        y -= 18
        c.drawString(72, y, f"Generated: {datetime.now(timezone.utc).isoformat()}")
        y -= 36
        c.setFont("Helvetica-Bold", 14)
        c.drawString(72, y, "Category Totals")
        y -= 20
        c.setFont("Helvetica", 11)
        for row in _aggregate_by_category(db, version.id):
            c.drawString(72, y, f"{row['category']}: ${row['total']:,.0f}")
            y -= 16
            if y < 72:
                c.showPage()
                y = height - 72
        y -= 20
        c.setFont("Helvetica-Bold", 14)
        c.drawString(72, y, "Confidence")
        y -= 20
        c.setFont("Helvetica", 11)
        c.drawString(
            72,
            y,
            f"High: {version.high_confidence_count}  "
            f"Medium: {version.medium_confidence_count}  "
            f"Low: {version.low_confidence_count}",
        )
        footnotes = _exog_footnotes(db, version.id)
        if footnotes:
            y -= 28
            if y < 120:
                c.showPage()
                y = height - 72
            c.setFont("Helvetica-Bold", 14)
            c.drawString(72, y, "Driver-dependent figures")
            y -= 18
            c.setFont("Helvetica", 10)
            c.drawString(
                72,
                y,
                "Exogenous drivers underpin these lines; kill-switch does not un-publish.",
            )
            y -= 16
            for line in footnotes[:40]:
                if y < 72:
                    c.showPage()
                    y = height - 72
                c.drawString(72, y, line[:110])
                y -= 14
        c.save()
        return buf.getvalue()
    except ImportError:
        lines = [
            f"Board Pack: {version.name}",
            f"Status: {version.status}",
            f"Horizon: {version.horizon_months}",
            "",
            "Category Totals:",
        ]
        for row in _aggregate_by_category(db, version.id):
            lines.append(f"  {row['category']}: ${row['total']:,.0f}")
        footnotes = _exog_footnotes(db, version.id)
        if footnotes:
            lines.append("")
            lines.append("Driver-dependent figures:")
            lines.extend(f"  {f}" for f in footnotes)
        return "\n".join(lines).encode("utf-8")


def save_board_pack(
    db: Session,
    version: ForecastVersion,
    fmt: str,
    output_dir: str,
) -> Path:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    safe = version.name.replace("/", "-").replace(" ", "_")
    if fmt == "pptx":
        data = build_board_pack_pptx(db, version)
        path = output / f"{safe}_board_pack.pptx"
    else:
        data = build_board_pack_pdf(db, version)
        path = output / f"{safe}_board_pack.pdf"
    path.write_bytes(data)
    return path
