"""Document processing pipeline: parse, chunk, embed, and store.

Leverages LangChain document loaders for parsing and RecursiveCharacterTextSplitter
for chunking, then stores embeddings in ChromaDB via the vector_store module.
"""

import logging
import os
import uuid
from pathlib import Path
from typing import Any

from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy.orm import Session

from app.config import settings
from app.models.document import Document, DocumentChunk
from app.services import vector_store

logger = logging.getLogger(__name__)

_splitter = RecursiveCharacterTextSplitter(
    separators=["\n\n", "\n", ". ", " "],
    chunk_size=settings.context_chunk_size,
    chunk_overlap=settings.context_chunk_overlap,
    length_function=len,
)

SUPPORTED_TYPES = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".doc": "docx",
    ".pptx": "pptx",
    ".xlsx": "xlsx",
    ".xls": "xlsx",
    ".csv": "csv",
    ".txt": "txt",
    ".md": "txt",
    ".html": "html",
    ".htm": "html",
}


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    return SUPPORTED_TYPES.get(ext, "txt")


def _load_pdf(file_path: str) -> list[dict]:
    """Load PDF using pypdf (via LangChain loader)."""
    try:
        from langchain_community.document_loaders import PyPDFLoader
        loader = PyPDFLoader(file_path)
        docs = loader.load()
        return [{"content": d.page_content, "metadata": d.metadata} for d in docs]
    except Exception as e:
        logger.warning(f"PyPDF failed, trying pdfplumber: {e}")
        import pdfplumber
        pages = []
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append({"content": text, "metadata": {"page": i + 1, "source": file_path}})
        return pages


def _load_docx(file_path: str) -> list[dict]:
    from langchain_community.document_loaders import Docx2txtLoader
    loader = Docx2txtLoader(file_path)
    docs = loader.load()
    return [{"content": d.page_content, "metadata": d.metadata} for d in docs]


def _load_pptx(file_path: str) -> list[dict]:
    from pptx import Presentation
    prs = Presentation(file_path)
    slides = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)
        if texts:
            slides.append({
                "content": "\n".join(texts),
                "metadata": {"slide": i + 1, "source": file_path},
            })
    return slides


def _load_xlsx(file_path: str) -> list[dict]:
    import pandas as pd
    pages = []
    try:
        xls = pd.ExcelFile(file_path)
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name)
            text = df.to_string(index=False)
            if text.strip():
                pages.append({
                    "content": text,
                    "metadata": {"sheet": sheet_name, "source": file_path},
                })
    except Exception:
        from langchain_community.document_loaders import CSVLoader
        loader = CSVLoader(file_path)
        docs = loader.load()
        pages = [{"content": d.page_content, "metadata": d.metadata} for d in docs]
    return pages


def _load_csv(file_path: str) -> list[dict]:
    from langchain_community.document_loaders import CSVLoader
    loader = CSVLoader(file_path)
    docs = loader.load()
    return [{"content": d.page_content, "metadata": d.metadata} for d in docs]


def _load_text(file_path: str) -> list[dict]:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
    return [{"content": content, "metadata": {"source": file_path}}]


def _load_html(file_path: str) -> list[dict]:
    from bs4 import BeautifulSoup
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    text = soup.get_text(separator="\n", strip=True)
    return [{"content": text, "metadata": {"source": file_path}}]


def load_url(url: str) -> list[dict]:
    """Fetch and parse a URL into document pages."""
    import httpx
    from bs4 import BeautifulSoup

    resp = httpx.get(url, follow_redirects=True, timeout=30)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    text = soup.get_text(separator="\n", strip=True)
    title = soup.title.string if soup.title else url
    return [{"content": text, "metadata": {"source": url, "title": title}}]


_LOADERS = {
    "pdf": _load_pdf,
    "docx": _load_docx,
    "pptx": _load_pptx,
    "xlsx": _load_xlsx,
    "csv": _load_csv,
    "txt": _load_text,
    "html": _load_html,
}


def parse_file(file_path: str, file_type: str) -> list[dict]:
    """Parse a file into a list of {content, metadata} dicts."""
    loader = _LOADERS.get(file_type, _load_text)
    return loader(file_path)


def chunk_pages(pages: list[dict]) -> list[dict]:
    """Split parsed pages into smaller chunks suitable for embedding."""
    all_chunks = []
    for page in pages:
        text = page["content"]
        if not text.strip():
            continue
        splits = _splitter.split_text(text)
        for split_text in splits:
            all_chunks.append({
                "content": split_text,
                "metadata": {**page.get("metadata", {}), "char_count": len(split_text)},
            })
    return all_chunks


def generate_summary(pages: list[dict], max_chars: int = 500) -> str:
    """Generate a brief summary from the first few pages of content."""
    combined = " ".join(p["content"][:300] for p in pages[:3])
    if len(combined) > max_chars:
        combined = combined[:max_chars] + "..."
    return combined


async def process_document(
    db: Session,
    doc: Document,
    file_path: str,
) -> Document:
    """Full processing pipeline: parse -> chunk -> embed -> store.
    
    Updates the Document record with status, chunk_count, etc.
    """
    try:
        logger.info(f"Processing document {doc.id}: {doc.original_name} ({doc.file_type})")

        if doc.file_type == "url" and doc.source_url:
            pages = load_url(doc.source_url)
        else:
            pages = parse_file(file_path, doc.file_type)

        if not pages:
            doc.status = "failed"
            doc.error_message = "No content could be extracted from the file."
            db.commit()
            return doc

        chunks = chunk_pages(pages)
        if not chunks:
            doc.status = "failed"
            doc.error_message = "File content could not be split into meaningful chunks."
            db.commit()
            return doc

        doc.description = generate_summary(pages)
        doc.total_chars = sum(len(c["content"]) for c in chunks)

        chunk_ids = []
        chunk_texts = []
        chunk_metadatas = []
        db_chunks = []

        for i, chunk_data in enumerate(chunks):
            chunk_id = str(uuid.uuid4())
            chunk_ids.append(chunk_id)
            chunk_texts.append(chunk_data["content"])
            chunk_metadatas.append({
                "document_id": doc.id,
                "user_id": doc.user_id,
                "scope": doc.scope,
                "file_type": doc.file_type,
                "original_name": doc.original_name,
                "chunk_index": i,
                **(chunk_data.get("metadata", {})),
            })
            db_chunks.append(DocumentChunk(
                id=chunk_id,
                document_id=doc.id,
                chunk_index=i,
                content=chunk_data["content"],
                chunk_metadata=chunk_data.get("metadata"),
                char_count=len(chunk_data["content"]),
                token_estimate=len(chunk_data["content"]) // 4,
                embedding_id=chunk_id,
            ))

        vector_store.add_chunks(doc.id, chunk_ids, chunk_texts, chunk_metadatas)

        db.add_all(db_chunks)
        doc.chunk_count = len(db_chunks)
        doc.status = "ready"
        db.commit()

        logger.info(
            f"Document {doc.id} processed: {len(db_chunks)} chunks, "
            f"{doc.total_chars} chars"
        )
        return doc

    except Exception as e:
        logger.error(f"Document processing failed for {doc.id}: {e}", exc_info=True)
        doc.status = "failed"
        doc.error_message = str(e)[:500]
        db.commit()
        return doc
