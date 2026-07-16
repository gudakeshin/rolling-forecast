"""CSV/Excel adapter for actuals data ingestion."""

import hashlib
import logging
from typing import Any

import pandas as pd

from app.services.ingestion.base import IActualsProvider, IngestionResult

logger = logging.getLogger(__name__)

# Expected column mappings (flexible naming)
COLUMN_ALIASES = {
    "account_code": ["account_code", "account_id", "gl_code", "code", "acct_code"],
    "account_name": ["account_name", "name", "description", "line_item", "gl_name"],
    "category": ["category", "type", "account_type", "pl_category"],
    "business_unit": ["business_unit", "bu", "department", "cost_center", "segment"],
    "geography": ["geography", "region", "geo", "country"],
    "product_line": ["product_line", "product", "product_group"],
    "period": ["period", "date", "month", "fiscal_period"],
    "value": ["value", "amount", "balance", "actual"],
    "currency": ["currency", "ccy", "curr"],
}


class CSVActualsProvider(IActualsProvider):
    """Adapter for ingesting actuals from CSV or Excel files."""

    @property
    def source_type(self) -> str:
        return "csv"

    async def pull_actuals(self, source_config: dict[str, Any]) -> IngestionResult:
        """Load and validate actuals from a CSV/Excel file (chunked reads)."""
        file_path = source_config.get("file_path", "")
        chunk_size = int(source_config.get("chunk_size", 5000))

        try:
            chunks: list[pd.DataFrame] = []
            chunk_warnings: list[str] = []

            if file_path.endswith(".csv"):
                reader = pd.read_csv(file_path, chunksize=chunk_size)
                for i, chunk in enumerate(reader):
                    chunk = self._normalize_chunk(chunk)
                    w = await self.validate(chunk)
                    hard = [x for x in w if x.startswith("Missing required columns:")]
                    if hard and i == 0:
                        return IngestionResult(
                            success=False,
                            error=hard[0],
                            warnings=w,
                            metadata={"columns": list(chunk.columns)},
                        )
                    chunk_warnings.extend(
                        f"chunk {i}: {x}" for x in w if not x.startswith("Missing required")
                    )
                    chunks.append(chunk)
            elif file_path.endswith((".xlsx", ".xls")):
                chunks = self._read_excel_chunked(file_path, chunk_size)
                if not chunks:
                    return IngestionResult(success=False, error="Excel file has no data rows")
                first = chunks[0]
                w = await self.validate(first)
                hard = [x for x in w if x.startswith("Missing required columns:")]
                if hard:
                    return IngestionResult(
                        success=False,
                        error=hard[0],
                        warnings=w,
                        metadata={"columns": list(first.columns)},
                    )
                for i, chunk in enumerate(chunks):
                    cw = await self.validate(chunk)
                    chunk_warnings.extend(
                        f"chunk {i}: {x}"
                        for x in cw
                        if not x.startswith("Missing required")
                    )
            else:
                return IngestionResult(
                    success=False, error=f"Unsupported file format: {file_path}"
                )

            if not chunks:
                return IngestionResult(success=False, error="File contains no rows")

            df = pd.concat(chunks, ignore_index=True)

            # Ensure period format is YYYY-MM
            df["period"] = df["period"].apply(self._normalize_period)

            # Fill missing currency
            if "currency" not in df.columns:
                df["currency"] = "USD"

            # Fill optional columns
            for col in ["geography", "product_line"]:
                if col not in df.columns:
                    df[col] = None

            file_hash = self._compute_hash(df)
            warnings = await self.validate(df)
            warnings = list(dict.fromkeys(warnings + chunk_warnings))  # dedupe, keep order

            periods = sorted(df["period"].unique())
            all_expected_periods = self._get_expected_periods(periods[0], periods[-1])
            missing = [p for p in all_expected_periods if p not in periods]
            completeness = (
                (len(periods) / len(all_expected_periods) * 100) if all_expected_periods else 100.0
            )

            return IngestionResult(
                success=True,
                dataframe=df,
                row_count=len(df),
                period_start=periods[0],
                period_end=periods[-1],
                periods_count=len(periods),
                missing_periods=missing,
                completeness_pct=round(completeness, 1),
                file_hash=file_hash,
                warnings=warnings,
                metadata={
                    "columns": list(df.columns),
                    "unique_accounts": df["account_code"].nunique(),
                    "unique_bus": df["business_unit"].nunique() if "business_unit" in df.columns else 0,
                    "categories": df["category"].unique().tolist() if "category" in df.columns else [],
                    "chunks_read": len(chunks),
                },
            )

        except Exception as e:
            logger.error(f"CSV ingestion failed: {e}", exc_info=True)
            return IngestionResult(success=False, error=str(e))

    def _normalize_chunk(self, df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        df.columns = [c.strip().lower().replace(" ", "_") for c in df.columns]
        return self._map_columns(df)

    def _read_excel_chunked(self, file_path: str, chunk_size: int) -> list[pd.DataFrame]:
        """Read Excel with openpyxl read_only, yielding DataFrame chunks."""
        try:
            from openpyxl import load_workbook
        except ImportError:
            df = pd.read_excel(file_path)
            return [self._normalize_chunk(df)] if len(df) else []

        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            wb.close()
            return []
        columns = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(header)]

        chunks: list[pd.DataFrame] = []
        buf: list[tuple] = []
        for row in rows_iter:
            if all(v is None for v in row):
                continue
            buf.append(row)
            if len(buf) >= chunk_size:
                chunk = pd.DataFrame(buf, columns=columns)
                chunks.append(self._normalize_chunk(chunk))
                buf = []
        if buf:
            chunk = pd.DataFrame(buf, columns=columns)
            chunks.append(self._normalize_chunk(chunk))
        wb.close()
        return chunks

    async def validate(self, df: pd.DataFrame) -> list[str]:
        """Validate the DataFrame for data quality issues.

        Missing required columns are reported with a 'Missing required columns:'
        prefix so pull_actuals can promote them to hard rejection.
        """
        warnings = []

        # Check required columns (hard rejection upstream)
        required = ["account_code", "account_name", "period", "value"]
        missing_cols = [c for c in required if c not in df.columns]
        if missing_cols:
            warnings.append(f"Missing required columns: {missing_cols}")

        # Check for null values in critical columns
        for col in ["account_code", "period", "value"]:
            if col in df.columns:
                null_count = df[col].isnull().sum()
                if null_count > 0:
                    warnings.append(f"Column '{col}' has {null_count} null values")

        # Check for unexpected zero balances
        if "value" in df.columns:
            zero_accounts = (
                df.groupby("account_code")["value"]
                .apply(lambda x: (x == 0).all())
            )
            all_zero = zero_accounts[zero_accounts].index.tolist()
            if all_zero:
                warnings.append(
                    f"{len(all_zero)} accounts have all-zero balances: {all_zero[:5]}{'...' if len(all_zero) > 5 else ''}"
                )

        # Check data freshness
        if "period" in df.columns:
            latest_period = df["period"].max()
            warnings.append(f"Most recent actuals period: {latest_period}")

        return warnings

    def _map_columns(self, df: pd.DataFrame) -> pd.DataFrame:
        """Map flexible column names to standard names."""
        rename_map = {}
        for standard_name, aliases in COLUMN_ALIASES.items():
            for alias in aliases:
                if alias in df.columns and standard_name not in df.columns:
                    rename_map[alias] = standard_name
                    break
        if rename_map:
            df = df.rename(columns=rename_map)
        return df

    def _normalize_period(self, period_val) -> str:
        """Normalize period formats to YYYY-MM or pass through FY2026-P01 labels."""
        period_str = str(period_val).strip()
        if period_str.upper().startswith("FY") and "-P" in period_str.upper():
            # Canonicalize FY2026-P1 → FY2026-P01
            body = period_str[2:]
            fy_s, p_s = body.upper().replace("P", "").split("-", 1)
            return f"FY{int(fy_s)}-P{int(p_s):02d}"
        for fmt in ["%Y-%m", "%Y/%m", "%m/%Y", "%Y-%m-%d", "%m/%d/%Y"]:
            try:
                dt = pd.to_datetime(period_str, format=fmt)
                return dt.strftime("%Y-%m")
            except (ValueError, TypeError):
                continue
        try:
            dt = pd.to_datetime(period_str)
            return dt.strftime("%Y-%m")
        except Exception:
            return period_str

    def _compute_hash(self, df: pd.DataFrame) -> str:
        """Compute SHA-256 hash of the DataFrame for lineage tracking."""
        content = df.to_csv(index=False).encode("utf-8")
        return hashlib.sha256(content).hexdigest()

    def _get_expected_periods(self, start: str, end: str) -> list[str]:
        """Generate expected periods between start and end (calendar-aware)."""
        from app.services.period_calendar import get_calendar_config, period_range

        try:
            return period_range(start, end, get_calendar_config())
        except Exception:
            try:
                date_range = pd.date_range(start=start + "-01", end=end + "-01", freq="MS")
                return [d.strftime("%Y-%m") for d in date_range]
            except Exception:
                return []
