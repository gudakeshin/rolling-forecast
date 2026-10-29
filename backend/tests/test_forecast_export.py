"""Whole-version forecast export (plan 4.3): every row, scoped, safe to open."""

from __future__ import annotations

import csv
import io
import os
import uuid

import pytest
from fastapi.testclient import TestClient
from passlib.context import CryptContext

_pwd = CryptContext(schemes=["bcrypt"], deprecated="auto")
N_LINES = 30
N_PERIODS = 24  # 720 rows: well past any single page


@pytest.fixture
def client():
    os.environ["SEED_DEMO_USERS"] = "true"
    os.environ["APP_ENV"] = "test"
    from app.main import app
    from app.rate_limit import limiter

    try:
        limiter.reset()
    except Exception:
        pass
    with TestClient(app) as c:
        yield c


@pytest.fixture
def data():
    from app.database import SessionLocal
    from app.models.business_unit import BusinessUnit
    from app.models.company import Company
    from app.models.forecast import ForecastLineResult, ForecastVersion
    from app.models.line_item import LineItem
    from app.models.user import Role, User

    tag = uuid.uuid4().hex[:8]
    db = SessionLocal()
    try:
        role = db.query(Role).filter(Role.name == "analyst").first()
        bus = {}
        for side in ("A", "B"):
            co = Company(name=f"ExportCo-{side}-{tag}")
            db.add(co)
            db.flush()
            bu = BusinessUnit(name=f"ExportBU-{side}-{tag}", company_id=co.id)
            db.add(bu)
            db.flush()
            bus[side] = bu.id
        version = ForecastVersion(name=f"Export v1 {tag}", business_unit_id=bus["A"])
        db.add(version)
        db.flush()
        for i in range(N_LINES):
            li = LineItem(
                account_code=f"EX-{tag}-{i:02d}",
                name="=HYPERLINK(\"http://x\")" if i == 0 else f"Line {i}",
                category="COGS" if i % 2 else "Revenue",
                business_unit_id=bus["A"],
            )
            db.add(li)
            db.flush()
            for p in range(N_PERIODS):
                db.add(ForecastLineResult(
                    version_id=version.id, line_item_id=li.id,
                    period=f"{2026 + p // 12}-{p % 12 + 1:02d}", p50=100.0 + p,
                    is_overridden=(i == 1 and p == 0), override_value=(5.0 if i == 1 and p == 0 else None),
                ))
        users = {}
        for side in ("A", "B"):
            username = f"export_{side.lower()}_{tag}"
            db.add(User(
                email=f"{username}@test.local", username=username,
                hashed_password=_pwd.hash(username), full_name=username,
                business_unit_id=bus[side], role_id=role.id,
            ))
            users[side] = username
        db.commit()
        return {"version": version.id, "users": users}
    finally:
        db.close()


def _h(client, username):
    r = client.post("/api/auth/login", json={"username": username, "password": username})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _csv(client, data, **params):
    r = client.get(
        f"/api/panel/forecast-table/{data['version']}/export",
        params=params, headers=_h(client, data["users"]["A"]),
    )
    assert r.status_code == 200, r.text
    return list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig")))), r


def test_csv_has_every_row_not_a_page(client, data):
    rows, r = _csv(client, data)
    assert len(rows) == N_LINES * N_PERIODS
    assert r.headers["content-disposition"].endswith('.csv"')
    overridden = next(row for row in rows if row["Overridden"] == "True")
    assert overridden["Published value"] == "5.0"


def test_filters_apply(client, data):
    rows, _ = _csv(client, data, category="Revenue")
    assert len(rows) == (N_LINES // 2) * N_PERIODS
    assert {row["Category"] for row in rows} == {"Revenue"}


def test_formula_like_names_are_neutralized(client, data):
    rows, _ = _csv(client, data)
    assert {row["Line item"] for row in rows if "HYPERLINK" in row["Line item"]} == {
        "'=HYPERLINK(\"http://x\")"
    }


def test_xlsx_opens_with_all_rows(client, data):
    from openpyxl import load_workbook

    r = client.get(
        f"/api/panel/forecast-table/{data['version']}/export",
        params={"format": "xlsx"}, headers=_h(client, data["users"]["A"]),
    )
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == N_LINES * N_PERIODS + 1
    assert ws.freeze_panes == "C2"


def test_other_company_gets_404_and_bad_format_422(client, data):
    url = f"/api/panel/forecast-table/{data['version']}/export"
    assert client.get(url, headers=_h(client, data["users"]["B"])).status_code == 404
    assert client.get(url, params={"format": "pdf"}, headers=_h(client, data["users"]["A"])).status_code == 422
